"""Inline AI edit of one bullet (⌘K): evidence rules hold, uncited policy claims are refused (offline, mock LLM)."""
from __future__ import annotations

import pytest

from app.models.pitch import Pitch, Slide, SlideBullet
from app.models.policy import CoverageStatus, SourceRef
from app.pitch.evidence_pack import EvidenceItem, EvidencePack
from app.pitch.rewrite import RewriteOut, RewriteRefused, rewrite_bullet
from app.services.llm import LLMService, LLMUnavailable

POL = "hdfc_optima_secure_plus"


def _pack() -> EvidencePack:
    src = SourceRef(policy_id=POL, chunk_id="c-maternity", page=14, source_text="Maternity is covered after a 36-month waiting period.")
    item = EvidenceItem(evidence_id="E1", policy_id=POL, policy_name="HDFC", feature_key="maternity", feature_label="Maternity", status=CoverageStatus.CONDITIONAL, statement="Maternity covered after 36 months.", conditions=["36-month waiting period"], sources=[src])
    return EvidencePack(recommended_policy_id=POL, items=[item])


def _pitch() -> Pitch:
    b = SlideBullet(text="Maternity is covered subject to a 36-month waiting period.", source_chunk_ids=["c-maternity"], policy_id=POL, kind="policy")
    return Pitch(pitch_id="p", company_name="Acme", recommended_policy_id=POL, slides=[Slide(slide_number=2, title="Key benefits", bullets=[b])])


def test_rewrite_keeps_evidence_and_maps_ids_to_chunks():
    seen: dict[str, str] = {}

    def handler(system, user, schema):
        seen["user"] = user
        return RewriteOut(text="Maternity cover applies after a 36-month waiting period.", kind="policy", evidence_ids=["E1"], note=None)

    pitch = _pitch()
    new, note = rewrite_bullet(pitch, _pack(), 2, pitch.slides[0].bullets[0], "make it shorter", llm=LLMService(mock_handler=handler))
    assert new.source_chunk_ids == ["c-maternity"] and new.policy_id == POL and new.kind == "policy" and note is None
    assert "currently cites E1" in seen["user"] and "ADVISOR INSTRUCTION:\nmake it shorter" in seen["user"]


def test_rewrite_refuses_uncited_policy_claim():
    def handler(system, user, schema):
        return RewriteOut(text="Dental implants are covered up to INR 5 lakh.", kind="policy", evidence_ids=["E-nope"], note="The pack has no dental evidence.")

    pitch = _pitch()
    with pytest.raises(RewriteRefused, match="dental"):
        rewrite_bullet(pitch, _pack(), 2, pitch.slides[0].bullets[0], "add dental cover", llm=LLMService(mock_handler=handler))


def test_rewrite_downgrade_to_recommendation_drops_policy_id():
    def handler(system, user, schema):
        return RewriteOut(text="We recommend confirming the maternity waiting period with the insurer.", kind="recommendation", evidence_ids=[], note=None)

    pitch = _pitch()
    new, _ = rewrite_bullet(pitch, _pack(), 2, pitch.slides[0].bullets[0], "turn this into a next step", llm=LLMService(mock_handler=handler))
    assert new.kind == "recommendation" and new.source_chunk_ids == [] and new.policy_id is None


def test_rewrite_without_llm_is_unavailable_not_silent():
    pitch = _pitch()
    with pytest.raises(LLMUnavailable):
        rewrite_bullet(pitch, _pack(), 2, pitch.slides[0].bullets[0], "shorter", llm=LLMService(mock_handler=None))


def test_deck_links_brochure_pages_and_web_sources(tmp_path):
    """Every reference in the exported deck is a hyperlink: brochure citations open the PDF at the page, company facts open the web page."""
    from pptx import Presentation

    from app.models.pitch import Pitch, Slide, SlideBullet
    from app.models.policy import SourceRef
    from app.pitch.pptx_builder import build_pitch_deck

    ref = SourceRef(policy_id=POL, policy_name="Acme Secure", chunk_id="c1", page=7, section="Waiting Periods", source_text="Maternity: 36 months waiting period.")
    pitch = Pitch(pitch_id="p1", company_name="Acme", recommended_policy_id=POL, slides=[
        Slide(slide_number=1, title="Acme: what we understand", bullets=[SlideBullet(text="Acme has 12,000 employees.", kind="company", source_urls=["https://en.wikipedia.org/wiki/Acme_Corporation"])]),
        Slide(slide_number=2, title="Why Acme Secure fits", bullets=[SlideBullet(text="Maternity is covered after 36 months.", kind="policy", source_chunk_ids=["c1"], policy_id=POL)]),
        Slide(slide_number=3, title="Next steps", bullets=[SlideBullet(text="Confirm group terms.", kind="recommendation")]),
    ])
    out = build_pitch_deck(pitch, {"c1": ref}, tmp_path / "deck.pptx")
    prs = Presentation(str(out))
    links = {i: [r.hyperlink.address for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address] for i, s in enumerate(prs.slides)}
    texts = {i: " ".join(r.text for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs) for i, s in enumerate(prs.slides)}
    # slide 1 (company): marker + reference both point at the web page; the reference reads as host · path
    assert links[1] == ["https://en.wikipedia.org/wiki/Acme_Corporation"] * 2 and "en.wikipedia.org · wiki/Acme Corporation" in texts[1]
    # slide 2 (policy): marker + reference open the brochure PDF at the cited page through the API
    assert links[2] == [f"http://localhost:8000/api/policies/{POL}/document#page=7"] * 2 and "Acme Secure brochure, p.7, Waiting Periods" in texts[2]
    assert links[3] == [] and "Sources" not in texts[3]  # nothing cited, no reference block
